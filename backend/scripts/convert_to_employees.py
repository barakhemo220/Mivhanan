import openpyxl
import json
import re

wb = openpyxl.load_workbook(r"c:\Users\ENGENER\Desktop\Mivhanan\רשימת מכשירנים.xlsx", data_only=True)
sheet = wb.active

dept_map = {}
for rng in sheet.merged_cells.ranges:
    if rng.min_col <= 3 <= rng.max_col:
        top_val = sheet.cell(row=rng.min_row, column=rng.min_col).value
        if top_val:
            for r in range(rng.min_row, rng.max_row + 1):
                dept_map[r] = top_val

if 2 not in dept_map and 3 in dept_map:
    dept_map[2] = dept_map[3]

def normalize_dept(raw):
    if not raw:
        return "מכשיריית כללית"
    text = str(raw).replace('\n', ' ').strip()
    text = re.sub(r'ח\.?\s*גלם', 'חומר גלם', text)
    text = re.sub(r'ת\.?\s*כוח', 'תחנת כוח', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def get_base_dept_name(full_dept):
    clean = re.sub(r'^מכשיריית\s*', '', full_dept).strip()
    return clean

employees = []

for r in range(2, sheet.max_row + 1):
    c1 = sheet.cell(row=r, column=1)
    c2 = sheet.cell(row=r, column=2)
    c3 = sheet.cell(row=r, column=3)

    id_val = c1.value
    name_val = c2.value
    raw_dept = dept_map.get(r, c3.value)

    if not id_val and not name_val:
        continue

    is_bold = bool(c2.font and c2.font.bold)
    id_str = str(int(id_val)) if isinstance(id_val, (int, float)) else str(id_val).strip()
    name_str = str(name_val).replace('\n', ' ').strip() if name_val else ''

    # Handle row 19 specifically: "השם שצריך להיות הוא מור עסל"
    if id_str == '132668':
        name_str = 'מור עסל'

    full_dept = normalize_dept(raw_dept)
    base_dept = get_base_dept_name(full_dept)

    if is_bold:
        role = f"מנהל מכשיריית {base_dept}"
        level = 3
        is_admin = False
    else:
        role = f"מכשירן {base_dept}"
        level = 2
        is_admin = False

    emp_obj = {
        "id": id_str,
        "name": name_str,
        "department": full_dept,
        "role": role,
        "level": level,
        "isAdmin": is_admin,
        "email": f"{id_str}@company.com"
    }
    employees.append(emp_obj)

# Add System Administrator
employees.append({
    "id": "9999",
    "name": "מנהל מערכת",
    "department": "הנהלת מכשור ובקרה",
    "role": "מנהל מערכת ראשי (Admin)",
    "level": 3,
    "isAdmin": True,
    "email": "admin@company.com"
})

target_path = r"c:\Users\ENGENER\Desktop\Mivhanan\backend\data\employees.json"
with open(target_path, "w", encoding="utf-8") as f:
    json.dump(employees, f, ensure_ascii=False, indent=2)

print(f"Successfully wrote {len(employees)} employees to {target_path}")
